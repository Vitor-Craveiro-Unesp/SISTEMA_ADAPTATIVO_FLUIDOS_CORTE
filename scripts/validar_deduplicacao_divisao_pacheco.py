"""Verificação independente, somente leitura, das projeções deduplicadas."""

from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DIV = ROOT / "data" / "processed" / "divisao"
BEFORE = DIV / "backup_pre_deduplicacao"
AFTER = DIV if "--final" in sys.argv else DIV / ".dedup_stage"
SUFFIX = " [POSSÍVEL DUPLICATA]"
SPECS = {
    "desempenho_usinagem_v1.xlsx": (
        lambda x: x.startswith("comparativo | potencia | condicao_")
        or x.startswith("comparativo | vida_ferramenta_volume_removido | condicao_1 |"),
        lambda x: x.startswith("usinabilidade | potencia_corte | condicao_")
        or x.startswith("usinabilidade | vida_ferramenta_volume_removido | condicao_1 |")
        or x.startswith("comparativo | vida_ferramenta_volume_removido | somatorio |"),
        5,
    ),
    "propriedades_estabilidade_fluido_v1.xlsx": (
        lambda x: x.startswith("comparativo | nevoa |"),
        lambda x: False,
        1,
    ),
}


def base(value):
    return str(value or "").removesuffix(SUFFIX)


def style(cell):
    return (
        cell.fill.patternType,
        cell.fill.fgColor.type,
        cell.fill.fgColor.rgb,
        cell.font.bold,
        cell.font.color.type if cell.font.color else None,
        cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None,
        cell.number_format,
    )


for file, (drop, clear_flag, expected_drop) in SPECS.items():
    old = load_workbook(BEFORE / file, data_only=False)
    new = load_workbook(AFTER / file, data_only=False)
    assert old.sheetnames == new.sheetnames
    original = old["Leia-me"]["B3"].value
    summary = old["Leia-me"]["B4"].value
    kept_cols = [c for c in range(2, original + summary + 2)
                 if c > original + 1 or not drop(base(old["Dados"].cell(1, c).value))]
    assert (original + summary) - len(kept_cols) == expected_drop
    assert new["Leia-me"]["B3"].value == original - expected_drop
    assert new["Leia-me"]["B4"].value == summary

    for new_col, old_col in enumerate(kept_cols, start=2):
        old_header = old["Dados"].cell(1, old_col).value
        expected_header = base(old_header) if clear_flag(base(old_header)) else old_header
        assert new["Dados"].cell(1, new_col).value == expected_header, (file, old_col, "header")
        for row in range(2, 13):
            a = old["Dados"].cell(row, old_col)
            b = new["Dados"].cell(row, new_col)
            assert a.value == b.value, (file, old_col, row, "value")
            assert style(a) == style(b), (file, old_col, row, "style")
    last = 1 + len(kept_cols)
    assert all(new["Dados"].cell(1, c).value is None for c in range(last + 1, old["Dados"].max_column + 1))
    assert all(new["Dados"].cell(r, c).value is None
               for c in range(2 + original - expected_drop, 2 + original - expected_drop + summary)
               for r in range(2, 13)), (file, "summary filled")

    old_dict = [row for row in list(old["Dicionário"].values)[1:] if not drop(base(row[0]))]
    new_dict = list(new["Dicionário"].values)[1:1 + len(old_dict)]
    assert len(old_dict) == len(new_dict)
    for a, b in zip(old_dict, new_dict):
        assert a[:7] == b[:7], (file, "dictionary fields", a[0])
        assert a[9:11] == b[9:11], (file, "dictionary missing/status", a[0])
        if clear_flag(base(a[0])):
            assert b[7:9] == ("NÃO", None), (file, "dictionary flag", a[0])
        else:
            assert a[7:9] == b[7:9], (file, "dictionary flag changed", a[0])
            assert a[11] == b[11], (file, "dictionary note changed", a[0])
    assert all(x is None for row in list(new["Dicionário"].values)[1 + len(old_dict):] for x in row)

    old_trace = [row for row in list(old["Rastreabilidade"].values)[1:] if not drop(base(row[1]))]
    new_trace = list(new["Rastreabilidade"].values)[1:1 + len(old_trace)]
    assert len(old_trace) == len(new_trace) == 11 * (original - expected_drop)
    for a, b in zip(old_trace, new_trace):
        assert a[:6] == b[:6], (file, "trace value/source", a[0], a[1])
        assert a[8] == b[8], (file, "trace status", a[0], a[1])
        if clear_flag(base(a[1])):
            assert b[6:8] == ("NÃO", None), (file, "trace flag", a[0], a[1])
        else:
            assert a[6:8] == b[6:8], (file, "trace flag changed", a[0], a[1])
    assert all(x is None for row in list(new["Rastreabilidade"].values)[1 + len(old_trace):] for x in row)
    assert new["Dados"].freeze_panes == old["Dados"].freeze_panes
    assert new["Dicionário"].freeze_panes == old["Dicionário"].freeze_panes
    assert new["Rastreabilidade"].freeze_panes == old["Rastreabilidade"].freeze_panes
    print(file, "OK", expected_drop, "retiradas;", original - expected_drop, "originais;", len(new_trace), "registros de origem preservados")

for file in ("qualidade_peca_processo_v1.xlsx", "economico_v1.xlsx"):
    assert (DIV / file).read_bytes() == (BEFORE / file).read_bytes(), file
    print(file, "OK sem alteração")
