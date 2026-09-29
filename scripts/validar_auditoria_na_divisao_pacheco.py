"""Conferência independente da edição localizada do XLSX econômico."""

from hashlib import sha256
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DIVISAO = ROOT / "data" / "processed" / "divisao"
BACKUP = DIVISAO / "backup_pre_auditoria_na"
FILES = (
    "desempenho_usinagem_v1.xlsx",
    "qualidade_peca_processo_v1.xlsx",
    "propriedades_estabilidade_fluido_v1.xlsx",
    "economico_v1.xlsx",
)
RED = {"AA10": 2242.7, "I10": 0.8918, "N10": 117.5584,
       "AA11": 2450.7, "I11": 0.8161, "N11": 117.4828}
GREEN = {"E10", "F10", "G10", "K10", "L10", "M10", "E11", "F11", "M11"}
PINK = {"B10", "D10", "P10", "Q10", "U10", "B11", "D11", "P11", "Q11"}
ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}


def digest(path):
    return sha256(path.read_bytes()).hexdigest()


def rgb(cell):
    color = cell.fill.fgColor
    return color.rgb[-6:].upper() if color.type == "rgb" and isinstance(color.rgb, str) else None


def fmt(cell):
    return (cell._style, cell.number_format)


if (DIVISAO / "backup_pre_deduplicacao").exists():
    subprocess.run([sys.executable, str(ROOT / "scripts" / "validar_deduplicacao_divisao_pacheco.py"), "--final"], check=True)
else:
    for name in FILES[:3]:
        assert digest(DIVISAO / name) == digest(BACKUP / name), f"Arquivo alterado fora do econômico: {name}"

before = load_workbook(BACKUP / FILES[3], data_only=False)
after = load_workbook(DIVISAO / FILES[3], data_only=False)
assert before.sheetnames == after.sheetnames
value_changes = []
style_changes = []
formula_errors = []
for name in before.sheetnames:
    a, b = before[name], after[name]
    for row in b:
        for cell in row:
            prior = a[cell.coordinate]
            if prior.value != cell.value:
                value_changes.append((name, cell.coordinate, prior.value, cell.value))
            if fmt(prior) != fmt(cell):
                style_changes.append((name, cell.coordinate))
            if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in ERRORS):
                formula_errors.append((name, cell.coordinate, cell.value))

note_rows = (*range(16, 23), *range(24, 31))
expected_value_changes = {("Dados", addr) for addr in RED} | {("Leia-me", f"{col}{row}") for row in note_rows for col in "AB"}
assert {(name, addr) for name, addr, _, _ in value_changes} == expected_value_changes, value_changes
expected_style_changes = {("Dados", addr) for addr in RED.keys() | GREEN | PINK} | {("Leia-me", f"{col}{row}") for row in note_rows for col in "AB"}
assert set(style_changes) == expected_style_changes, f"Estilos não previstos: {set(style_changes) ^ expected_style_changes}"

data = after["Dados"]
for addr, value in RED.items():
    assert data[addr].value == value and rgb(data[addr]) == "FF0000", addr
for addr in GREEN:
    assert data[addr].value == 0 and rgb(data[addr]) == "00B050", addr
for addr in PINK:
    assert data[addr].value == "NA" and rgb(data[addr]) == "FFB6C1", addr
assert len(RED) == 6 and len(GREEN) == 9 and len(PINK) == 9
assert all(data[addr].value == "NA" for addr in ["B10", "B11", "D10", "D11", "Q10", "Q11"])
assert not formula_errors, formula_errors
readme = after["Leia-me"]
assert [readme[f"B{row}"].value for row in (20, 21, 22)] == [6, 9, 9]
print("OK: 6 vermelhas; 9 verdes; 9 rosas; 9 NA remanescentes; sem erros de fórmula.")
print(f"Mudanças de valor: {len(value_changes)} (6 dados + 28 células de notas).")
print(f"Mudanças de estilo na auditoria econômica: {len(style_changes)}; as demais planilhas foram validadas na versão atual.")
