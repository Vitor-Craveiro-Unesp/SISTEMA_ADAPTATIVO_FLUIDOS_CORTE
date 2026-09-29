"""Diagnóstico somente leitura das quatro visões da matriz candidata."""

from collections import Counter
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string


ROOT = Path(__file__).resolve().parents[1]
DIV = ROOT / "data" / "processed" / "divisao"
RAW = load_workbook(ROOT / "data" / "raw" / "ensaio_bancada_alunos (2).xlsx", data_only=True)
SOURCE = {
    "ANÁLISE USINABILIDADE": RAW.worksheets[3],
    "ANÁLISE EMULSÕES": RAW.worksheets[4],
    "COMPARATIVO": RAW.worksheets[5],
    "AVALIAÇÃO ECONÔMICA": RAW.worksheets[6],
}


def same(a, b):
    if a is None and b == "NA":
        return True
    try:
        return abs(float(a) - float(b)) < 1e-8
    except (TypeError, ValueError):
        return str(a) == str(b)


for path in sorted(DIV.glob("*.xlsx")):
    wb = load_workbook(path, data_only=False)
    data, dictionary, trace, readme = wb.worksheets
    original_count = readme["B3"].value
    summary_count = readme["B4"].value
    header_map = {
        re.sub(r" \[POSS.*DUPLICATA\]$", "", data.cell(1, c).value): c
        for c in range(2, original_count + 2)
    }
    row_map = {data.cell(r, 1).value: r for r in range(2, 13)}
    raw_match = Counter()
    comparable_shift = Counter()
    data_match = Counter()
    mismatches = []
    trace_address_mismatches = []
    active_trace = [row for row in list(trace.values)[1:] if row[0] is not None]
    for fluid, label, sheet, address, raw_value, *_ in active_trace:
        source = SOURCE[sheet]
        source_value = source[address].value
        col_letters, source_row = coordinate_from_string(address)
        source_value_plus6 = source[f"{col_letters}{source_row + 6}"].value
        if sheet == "COMPARATIVO":
            source_status = "direct_coincidence" if same(source_value, raw_value) else "plus6" if same(source_value_plus6, raw_value) else "neither"
        elif same(source_value, raw_value):
            source_status = "direct"
        elif raw_value == "NA" and source[address].data_type == "e":
            source_status = "formula_error_cache_mapped_NA"
        else:
            source_status = "neither"
        raw_match[(sheet, source_status)] += 1
        if sheet != "COMPARATIVO" and not same(source_value, raw_value):
            trace_address_mismatches.append((fluid, label, sheet, address, raw_value, source_value))
        if sheet == "COMPARATIVO":
            comparable_shift["plus6_matches" if same(source_value_plus6, raw_value) else "plus6_fails"] += 1
        col, row = header_map[label], row_map[fluid]
        value = data.cell(row, col).value
        key = "equal" if same(value, raw_value) else "changed"
        data_match[key] += 1
        if key == "changed":
            mismatches.append((fluid, data.cell(row, col).coordinate, raw_value, value))
    summary_cells = [data.cell(r, c).value for r in range(2, 13)
                     for c in range(original_count + 2, original_count + summary_count + 2)]
    print(path.name)
    print(" originais", original_count, "rastreabilidade", len(active_trace),
          "resumos em branco", sum(v is None for v in summary_cells), "de", len(summary_cells))
    print(" fonte", dict(raw_match))
    if comparable_shift:
        print(" comparativo_offset", dict(comparable_shift))
    print(" dados vs rastreabilidade", dict(data_match), "diferenças", mismatches[:10])
    if trace_address_mismatches:
        print(" endereços não comparativo divergentes", trace_address_mismatches[:10])
