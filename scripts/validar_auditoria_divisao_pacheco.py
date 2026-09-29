"""Valida que a auditoria ampliada não mudou valores experimentais."""

from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DIR = ROOT / "data" / "processed" / "divisao"
BACKUP = DIR / "backup_pre_auditoria_todos"
EXPECTED = {
    "desempenho_usinagem_v1.xlsx": (12, 16, 23, 17, 22),
    "qualidade_peca_processo_v1.xlsx": (0, 16, 21, 17, 55),
    "propriedades_estabilidade_fluido_v1.xlsx": (11, 16, 24, 34, 363),
    "economico_v1.xlsx": (0, 24, 30, 26, 286),
}

if any(load_workbook(DIR / name, read_only=True)["Leia-me"]["B3"].value != spec[3]
       for name, spec in EXPECTED.items()):
    print("Estrutura posterior à auditoria original detectada; executando o validador de deduplicação atual.")
    subprocess.run([sys.executable, str(ROOT / "scripts" / "validar_deduplicacao_divisao_pacheco.py"), "--final"], check=True)
    raise SystemExit(0)


def color(cell):
    c = cell.fill.fgColor
    return c.rgb[-6:].upper() if c.type == "rgb" and isinstance(c.rgb, str) else None


for name, (expected_amber, first, last, original_fields, expected_blank) in EXPECTED.items():
    before = load_workbook(BACKUP / name, data_only=False)
    after = load_workbook(DIR / name, data_only=False)
    assert before.sheetnames == after.sheetnames, name
    data = after["Dados"]
    amber = [data.cell(r, c).coordinate for r in range(2, 13) for c in range(2, original_fields + 2)
             if color(data.cell(r, c)) == "FFC000"]
    assert len(amber) == expected_amber, (name, amber)
    summary = [data.cell(r, c).value for r in range(2, 13)
               for c in range(original_fields + 2, data.max_column + 1)]
    assert len(summary) == expected_blank and all(v is None for v in summary), name
    style_changes = []
    for sheet_name in before.sheetnames:
        prior, current = before[sheet_name], after[sheet_name]
        for row in current:
            for cell in row:
                old = prior[cell.coordinate]
                if sheet_name != "Leia-me":
                    assert old.value == cell.value, (name, sheet_name, cell.coordinate, old.value, cell.value)
                    if old._style != cell._style:
                        style_changes.append((sheet_name, cell.coordinate))
                elif cell.row < first:
                    assert old.value == cell.value and old._style == cell._style, (name, sheet_name, cell.coordinate)
                else:
                    if cell.row > last:
                        assert old.value == cell.value and old._style == cell._style, (name, sheet_name, cell.coordinate)
                assert cell.data_type not in {"e", "f"}, (name, sheet_name, cell.coordinate, cell.value)
    assert set(style_changes) == {("Dados", address) for address in amber}, (name, style_changes)
    assert all(after["Leia-me"][f"A{row}"].value is not None for row in range(first, last + 1)), name
    print(f"{name}: {expected_amber} células laranja; {expected_blank} resumos vazios; dados e rastreabilidade preservados.")

economic = load_workbook(DIR / "economico_v1.xlsx", data_only=False)["Dados"]
red = {"AA10", "I10", "N10", "AA11", "I11", "N11"}
green = {"E10", "F10", "G10", "K10", "L10", "M10", "E11", "F11", "M11"}
pink = {"B10", "D10", "P10", "Q10", "U10", "B11", "D11", "P11", "Q11"}
assert all(color(economic[a]) == "FF0000" and isinstance(economic[a].value, (int, float)) for a in red)
assert all(color(economic[a]) == "00B050" and economic[a].value == 0 for a in green)
assert all(color(economic[a]) == "FFB6C1" and economic[a].value == "NA" for a in pink)
print("Econômico: 6 reconstruções vermelhas, 9 zeros verdes e 9 NA rosas preservados.")
