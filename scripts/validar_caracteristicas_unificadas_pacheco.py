"""Validação independente, somente leitura, da unificação Pacheco."""

from pathlib import Path
from math import isclose
import sys
import openpyxl


ROOT = Path.cwd()
DIV = ROOT / "data" / "processed" / "divisao"
RAW = ROOT / "data" / "raw" / "ensaio_bancada_alunos (2).xlsx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else DIV / "caracteristicas_unificadas_v1.xlsx"
PAIRS = [
    ("Desempenho usinagem", "desempenho_usinagem_v1.xlsx"),
    ("Qualidade peça processo", "qualidade_peca_processo_v1.xlsx"),
    ("Propriedades fluido", "propriedades_estabilidade_fluido_v1.xlsx"),
    ("Econômico", "economico_v1.xlsx"),
]
FLUIDS = list("ABCDEFGHJKM")


def same(a, b):
    if isinstance(a, (int, float)) and not isinstance(a, bool) and isinstance(b, (int, float)) and not isinstance(b, bool):
        return isclose(a, b, rel_tol=0, abs_tol=1e-10)
    if isinstance(a, (int, float)) and not isinstance(a, bool) and isinstance(b, str):
        try:
            return isclose(a, float(b.replace(",", ".")), rel_tol=0, abs_tol=1e-10)
        except ValueError:
            return False
    return a == b


new = openpyxl.load_workbook(OUT, data_only=True)
raw = openpyxl.load_workbook(RAW, data_only=True)
assert new.sheetnames == [x[0] for x in PAIRS] + ["Origem das colunas", "Legenda das cores"], new.sheetnames
total_values = 0
for tab, filename in PAIRS:
    src = openpyxl.load_workbook(DIV / filename, data_only=True)["Dados"]
    dst = new[tab]
    max_col = max(c.column for c in src[1] if c.value is not None)
    assert [dst.cell(r, 1).value for r in range(2, 13)] == FLUIDS, tab
    for row in range(1, 13):
        for col in range(1, max_col + 1):
            x, y = src.cell(row, col).value, dst.cell(row, col).value
            assert same(x, y), (tab, src.cell(row, col).coordinate, x, y)
            total_values += 1

lineage = new["Origem das colunas"]
assert lineage.max_row == sum(max(c.column for c in new[tab][1] if c.value is not None) for tab, _ in PAIRS) + 1
raw_verified = 0
numeric_text_conversion = 0
numeric_text_examples = []
skipped_reconstructions = 0
skipped_missing = 0
for row in lineage.iter_rows(min_row=2, max_col=12, values_only=True):
    tab, col_letter, _, raw_file, raw_sheet, raw_cols, addresses, filename, source_col, kind, _, flags = row
    assert filename == next(name for sheet, name in PAIRS if sheet == tab)
    assert col_letter == source_col
    if kind == "Identificador":
        continue
    if kind == "Resumo pendente":
        assert raw_file == "—" and raw_sheet == "—" and addresses == "—"
        continue
    assert raw_file == "data/raw/ensaio_bancada_alunos (2).xlsx", (tab, col_letter, raw_file)
    prefix = f"{raw_sheet}! "
    assert addresses.startswith(prefix), (tab, col_letter, addresses, prefix)
    parts = [part.split("=") for part in addresses[len(prefix):].split("; ")]
    assert [fluid for fluid, _ in parts] == FLUIDS, (tab, col_letter, parts)
    physical_cols = set()
    for i, (fluid, addr) in enumerate(parts, 2):
        assert raw_sheet in raw.sheetnames, (tab, raw_sheet)
        physical_cols.add("".join(ch for ch in addr if ch.isalpha()))
        value = new[tab][f"{col_letter}{i}"].value
        origin = raw[raw_sheet][addr].value
        if "azul-claro" in str(flags) and fluid in ("J", "K"):
            skipped_reconstructions += 1
            continue
        if value == "NA" or value is None:
            skipped_missing += 1
            continue
        assert same(value, origin), (tab, col_letter, fluid, addr, value, origin)
        if isinstance(value, (int, float)) and isinstance(origin, str):
            numeric_text_conversion += 1
            numeric_text_examples.append((tab, col_letter, fluid, addr, origin, value))
        raw_verified += 1
    assert set(raw_cols.split(", ")) == physical_cols, (tab, col_letter, raw_cols, physical_cols)

print(f"Abas: {len(new.sheetnames)}; células copiadas: {total_values}; colunas documentadas: {lineage.max_row - 1}")
print(f"Valores conferidos com a fonte física: {raw_verified}; reconstruções separadas: {skipped_reconstructions}; ausências/resumos: {skipped_missing}")
print(f"Números armazenados como texto na fonte e convertidos na projeção: {numeric_text_conversion}")
print(f"Conversões observadas: {numeric_text_examples}")

CATEGORICAL = {
    "Qualidade peça processo": ["J", "K", "L", "M", "N"],
    "Propriedades fluido": ["B", "D", "E", "G", "H", "J", "M", "T", "U", "V", "X", "Z", "AA", "AC", "AD", "AE"],
    "Econômico": ["R", "V"],
}
marked = 0
for tab, columns in CATEGORICAL.items():
    for col in columns:
        for row in range(2, 13):
            cell = new[tab][f"{col}{row}"]
            assert isinstance(cell.value, str) and cell.value.strip(), (tab, cell.coordinate, cell.value)
            assert cell.fill.fgColor.rgb == "FFDCC8F2", (tab, cell.coordinate, cell.fill.fgColor.rgb)
            marked += 1
assert marked == 253
for addr in ("U4", "U6"):
    cell = new["Propriedades fluido"][addr]
    assert cell.value == "*fungos"
    assert cell.border.bottom.color.rgb == "FFFFC000"
assert new["Legenda das cores"]["A1"].value == "Legenda das cores e variáveis categóricas"
print(f"Colunas categóricas sinalizadas: {sum(map(len, CATEGORICAL.values()))}; células lilás: {marked}; alertas combinados preservados: 2")
